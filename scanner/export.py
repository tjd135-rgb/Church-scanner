"""CSV + Excel export with a summary tab."""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import dor as dor_mod

log = logging.getLogger(__name__)


OUTPUT_COLUMNS = [
    "parcel_id", "situs_addr", "situs_city", "county",
    "owner_name", "owner_addr1", "owner_addr2",
    "owner_city", "owner_state", "owner_zip",
    "dor_uc", "dor_uc_description", "zoning",
    "lot_ac", "lot_sf", "bldg_sf",
    "land_val", "bldg_val", "total_val",
    "year_built", "sale_date", "sale_price",
    "land_utilization_ratio", "land_value_share", "hold_period_years",
    "score_acreage", "score_lur_inverse", "score_land_value_share",
    "score_hold_period", "redevelopment_score",
    "pass_land_use", "pass_buffer", "pass_address", "in_corridor_buffer",
    "source_layer",
]


def prepare_output(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "dor_uc" in out.columns:
        out["dor_uc_description"] = out["dor_uc"].map(dor_mod.describe)
    for c in OUTPUT_COLUMNS:
        if c not in out.columns:
            out[c] = pd.NA
    out = out[OUTPUT_COLUMNS]
    if "redevelopment_score" in out.columns:
        out = out.sort_values("redevelopment_score", ascending=False,
                              na_position="last").reset_index(drop=True)
    return out


def summary_lines(
    df: pd.DataFrame, scoring_docs: list[str],
    buffer_ft: float, cities: list[str], layer_used: str,
    validation_report: str,
) -> list[str]:
    lines = [
        "PARCEL SCAN SUMMARY — churches + private schools",
        "Corridor: US-1 (Federal Hwy) + Dixie Highway / SR 811",
        f"Cities: {', '.join(cities)}",
        f"Buffer: {buffer_ft:.0f} ft around road centerlines",
        f"Primary parcel layer: {layer_used}",
        f"Total parcels found: {len(df)}",
        "",
        "Breakdown by city:",
    ]
    if "situs_city" in df.columns and not df.empty:
        counts = df["situs_city"].fillna("(unknown)").astype(str).value_counts()
        for city, n in counts.items():
            lines.append(f"  {city}: {n}")
    lines.extend(["", "Breakdown by DOR use code:"])
    if "dor_uc" in df.columns and not df.empty:
        for code, n in df["dor_uc"].fillna(-1).astype("Int64").value_counts().items():
            lines.append(
                f"  {int(code) if code != -1 else 'unknown'} "
                f"({dor_mod.describe(int(code)) if code != -1 else 'unknown'}): {n}"
            )
    lines.extend(["", "Pass coverage (a parcel may hit multiple passes):"])
    for col, label in [
        ("pass_land_use", "land-use filter (DOR 71/72)"),
        ("pass_buffer", "spatial buffer"),
        ("pass_address", "situs-address regex"),
    ]:
        if col in df.columns:
            lines.append(f"  {label}: {int(df[col].fillna(False).astype(bool).sum())}")
    lines.extend(["", "Scoring model:", *[f"  {s}" for s in scoring_docs]])
    lines.extend(["", "Top 10 by redevelopment score:"])
    if not df.empty:
        top = df.head(10)[[
            "redevelopment_score", "parcel_id", "situs_addr", "situs_city",
            "owner_name", "lot_ac",
        ]]
        for _, r in top.iterrows():
            lines.append(
                f"  {r['redevelopment_score']:>5} | "
                f"{str(r['parcel_id']):>15} | "
                f"{str(r['situs_addr'])[:35]:35s} | "
                f"{str(r['situs_city'])[:18]:18s} | "
                f"{str(r['owner_name'])[:30]:30s} | "
                f"{r['lot_ac']:>6.2f} ac"
            )
    lines.extend(["", validation_report])
    return lines


def write_csv(df: pd.DataFrame, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    log.info("wrote %s (%d rows)", path, len(df))
    return path


def write_xlsx(
    df: pd.DataFrame, summary: list[str], path: Path,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    header_font = Font(bold=True, size=14)
    for i, line in enumerate(summary, start=1):
        cell = ws_sum.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = header_font
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws_sum.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Parcels")
    cols = list(df.columns)
    header_fill = PatternFill(start_color="FF204060", end_color="FF204060",
                              fill_type="solid")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, val in enumerate(row, start=1):
            if pd.isna(val):
                val = None
            elif hasattr(val, "isoformat"):
                val = val.isoformat()
            ws.cell(row=r_idx, column=c_idx, value=val)
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, start=1):
        width = min(max(len(str(c)), 10), 40)
        ws.column_dimensions[get_column_letter(i)].width = width

    ws_top = wb.create_sheet("Top 10")
    top_cols = [
        "redevelopment_score", "parcel_id", "situs_addr", "situs_city",
        "owner_name", "lot_ac", "land_utilization_ratio",
        "land_value_share", "hold_period_years",
        "land_val", "total_val",
    ]
    top_cols = [c for c in top_cols if c in df.columns]
    for i, c in enumerate(top_cols, start=1):
        cell = ws_top.cell(row=1, column=i, value=c)
        cell.font = Font(bold=True)
    for r_idx, row in enumerate(
        df.head(10)[top_cols].itertuples(index=False, name=None), start=2
    ):
        for c_idx, val in enumerate(row, start=1):
            if pd.isna(val):
                val = None
            elif hasattr(val, "isoformat"):
                val = val.isoformat()
            ws_top.cell(row=r_idx, column=c_idx, value=val)
    ws_top.freeze_panes = "A2"
    for i, _ in enumerate(top_cols, start=1):
        ws_top.column_dimensions[get_column_letter(i)].width = 22

    wb.save(path)
    log.info("wrote %s", path)
    return path
